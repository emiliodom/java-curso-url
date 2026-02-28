#!/usr/bin/env python3
"""Generate Sesion de Clase 6 and Guia de Aprendizaje 6 as .docx files."""
import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "..")

def set_cell_shading(cell, color_hex):
    from docx.oxml.ns import qn
    from lxml import etree
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shading = etree.SubElement(tcPr, qn('w:shd'))
    shading.set(qn('w:fill'), color_hex)
    shading.set(qn('w:val'), 'clear')

def add_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
        set_cell_shading(cell, '4F46E5')
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)
    # Data rows
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    return table

# ═══════════════════════════════════════════════════════════
# 1. SESION DE CLASE 6
# ═══════════════════════════════════════════════════════════
def create_sesion():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Universidad Rafael Landivar')
    r.bold = True; r.font.size = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Facultad de Humanidades').font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Departamento de Educacion').font.size = Pt(11)

    doc.add_paragraph()

    # Info table
    info_table = doc.add_table(rows=2, cols=2)
    info_table.style = 'Table Grid'
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.rows[0].cells[0].text = 'Profesorados con Especialidad en TIC'
    info_table.rows[0].cells[1].text = 'Nombre del curso: Introduccion al Desarrollo Web con Java'
    info_table.rows[1].cells[0].text = 'Ano Psicopedagogico'
    info_table.rows[1].cells[1].text = 'Numero de creditos: 4\nCiclo y modulo: Cuarto ciclo'
    for row in info_table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)

    doc.add_paragraph()

    # Title
    h = doc.add_heading('Secuencia de aprendizaje', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph('Sesion presencial/sincronica #6', style='Intense Quote')

    p = doc.add_paragraph()
    p.add_run('Tema: ').bold = True
    p.add_run('Persistencia en BD, Login con Spring Security, Roles, Renderizado condicional, Consultas JPA y Graficos con Chart.js')

    p = doc.add_paragraph()
    p.add_run('Enfoque del curso (perfil del estudiante): ').bold = True
    p.add_run('formacion de docentes TIC. Priorizamos comprension, capacidad de ensenar y evidencias por proceso.')

    doc.add_paragraph()

    # ── Informacion importante ──
    doc.add_heading('Informacion importante', level=2)

    add_styled_table(doc,
        ['Categoria', 'Detalle'],
        [
            ['Aprendizajes esperados',
             '- Implementa persistencia de datos en una base de datos sencilla (H2).\n'
             '- Desarrolla una funcionalidad con autenticacion/roles basicos (Spring Security).\n'
             '- Muestra u oculta elementos HTML segun el rol del usuario conectado.\n'
             '- Busca y filtra datos con metodos de Spring Data JPA.\n'
             '- Presenta datos en graficos con Chart.js.'],
            ['Contenidos tematicos',
             '- Spring Security: @Configuration, @EnableWebSecurity, SecurityFilterChain.\n'
             '- Usuarios en memoria con roles (ADMIN, USER).\n'
             '- Login personalizado con formulario Thymeleaf.\n'
             '- Renderizado condicional: sec:authorize, th:if.\n'
             '- Consultas JPA: findByTipo, findByTituloContainingIgnoreCase, countByTipo.\n'
             '- Graficos con Chart.js (barras, dona).'],
            ['Habilidades y destrezas',
             '- Configurar seguridad basica en una app Spring Boot.\n'
             '- Disenar vistas que cambian segun el rol.\n'
             '- Filtrar datos sin escribir SQL.\n'
             '- Conectar datos del backend con graficos en el frontend.'],
            ['Valores y actitudes',
             '- Responsabilidad en el manejo de datos sensibles.\n'
             '- Precision al definir permisos y accesos.'],
            ['Productos que evidencian el aprendizaje',
             'Login funcional + vista con renderizado condicional + filtro de datos + grafico con datos reales + explicacion didactica (5-8 parrafos).']
        ])

    doc.add_paragraph()

    # ── Desarrollo de la secuencia ──
    doc.add_heading('Desarrollo de la secuencia', level=2)

    # 1. Activacion
    doc.add_heading('1. Activacion de presaberes', level=3)
    p = doc.add_paragraph()
    p.add_run('Actividad: ').bold = True
    p.add_run('"Quien puede entrar?" — lluvia de ideas sobre por que proteger datos.')
    p = doc.add_paragraph()
    p.add_run('Tiempo estimado: ').bold = True
    p.add_run('10 minutos')
    p = doc.add_paragraph()
    p.add_run('Recursos: ').bold = True
    p.add_run('Pizarra o diapositiva con ejemplo: "Si cualquier persona puede ver los datos de los ninos del dispensario, que problema hay?"')
    p = doc.add_paragraph()
    p.add_run('Actividad que realizara el estudiante: ').bold = True
    p.add_run('explicar con sus palabras por que un sistema de informacion necesita login y que pasaria si no existiera.')

    # 2. Retroalimentacion
    doc.add_heading('2. Retroalimentacion', level=3)
    p = doc.add_paragraph()
    p.add_run('Actividad: ').bold = True
    p.add_run('repaso de CRUD funcional (Semana 5) y verificacion de que el proyecto compila.')
    p = doc.add_paragraph()
    p.add_run('Tiempo estimado: ').bold = True
    p.add_run('10 minutos')
    p = doc.add_paragraph()
    p.add_run('Recursos: ').bold = True
    p.add_run('Material HTML Sesion 6 (Material_Sesion_Clase_6_PERSISTENCIA/index.html).')
    p = doc.add_paragraph()
    p.add_run('Actividad que realizara el estudiante: ').bold = True
    p.add_run('verificar que su CRUD de Semana 5 funciona, hacer GET /recursos y POST para crear 1 registro.')

    # 3. Construccion
    doc.add_heading('3. Construccion del conocimiento (hands-on)', level=3)

    doc.add_heading('Actividad 1: Dependencias y login por defecto (10 min)', level=4)
    doc.add_paragraph('Agregar las 2 dependencias nuevas (spring-boot-starter-security y thymeleaf-extras-springsecurity6).')
    doc.add_paragraph('Ejecutar la app y observar que Spring bloquea todo con un login por defecto.')
    doc.add_paragraph('Evaluacion: el estudiante identifica la pantalla de login generada por Spring.')

    doc.add_heading('Actividad 2: Configuracion de seguridad (25 min)', level=4)
    doc.add_paragraph('Crear SeguridadConfig.java con SecurityFilterChain y usuarios en memoria.')
    doc.add_paragraph('Definir 2 usuarios: admin (ADMIN) y docente (USER).')
    doc.add_paragraph('Crear LoginControlador.java y templates/login.html.')
    doc.add_paragraph('Evaluacion: login funciona con ambos usuarios; credencial incorrecta muestra error.')

    doc.add_heading('Actividad 3: Roles y renderizado condicional (25 min)', level=4)
    doc.add_paragraph('Agregar namespace sec: a las vistas HTML.')
    doc.add_paragraph('Mostrar nombre del usuario con sec:authentication="name".')
    doc.add_paragraph('Ocultar boton "Borrar" para USER, mostrar solo para ADMIN.')
    doc.add_paragraph('Agregar enlace a "Estadisticas" visible solo para ADMIN.')
    doc.add_paragraph('Evaluacion: iniciar como admin y como docente; verificar que ven cosas diferentes.')

    doc.add_heading('Actividad 4: Consultas y filtrado (20 min)', level=4)
    doc.add_paragraph('Agregar metodos findByTipo y findByTituloContainingIgnoreCase al repositorio.')
    doc.add_paragraph('Modificar controlador para aceptar @RequestParam opcionales.')
    doc.add_paragraph('Agregar formulario de busqueda en la vista.')
    doc.add_paragraph('Evaluacion: filtrar por tipo "Video" y buscar por "java" funcionan correctamente.')

    doc.add_heading('Actividad 5: Graficos con Chart.js (20 min)', level=4)
    doc.add_paragraph('Agregar countByTipo al repositorio.')
    doc.add_paragraph('Crear ruta GET /recursos/estadisticas en el controlador.')
    doc.add_paragraph('Crear vista estadisticas.html con Chart.js (barras y dona).')
    doc.add_paragraph('Evaluacion: el grafico muestra datos reales de la BD.')

    # 4. Aplicacion
    doc.add_heading('4. Aplicacion del conocimiento', level=3)
    p = doc.add_paragraph()
    p.add_run('Actividad: Mini-reto con roles (15 min)').bold = True
    doc.add_paragraph('El estudiante agrega un tercer rol EDITOR que puede crear y editar pero no borrar.')
    doc.add_paragraph('Debe modificar SeguridadConfig y las vistas con sec:authorize.')
    doc.add_paragraph('Evaluacion: el nuevo rol funciona y la vista cambia correctamente.')

    # 5. Reflexion
    doc.add_heading('5. Reflexion sobre mi proceso de aprendizaje', level=3)
    p = doc.add_paragraph()
    p.add_run('Actividad: ').bold = True
    p.add_run('"Como explicaria esto a un estudiante de secundaria".')
    p = doc.add_paragraph()
    p.add_run('Tiempo estimado: ').bold = True
    p.add_run('10 minutos')
    doc.add_paragraph('El estudiante redacta 3 analogias: una para login, una para roles y una para consultas JPA.')

    # 6. Introduccion siguiente tema
    doc.add_heading('6. Introduccion al siguiente tema', level=3)
    p = doc.add_paragraph()
    p.add_run('Actividad: ').bold = True
    p.add_run('adelanto Semana 7 (deploy — publicar la app en internet).')
    p = doc.add_paragraph()
    p.add_run('Tiempo estimado: ').bold = True
    p.add_run('5 minutos')
    doc.add_paragraph('El estudiante anota que necesitaria preparar para llevar su app a produccion.')

    # 7. Cierre
    doc.add_heading('7. Actividad de cierre y despedida', level=3)
    p = doc.add_paragraph()
    p.add_run('Actividad: ').bold = True
    p.add_run('ticket de salida.')
    p = doc.add_paragraph()
    p.add_run('Tiempo estimado: ').bold = True
    p.add_run('5 minutos')
    doc.add_paragraph('El estudiante escribe una micro-guia de 5 pasos para agregar login a un proyecto Spring Boot existente.')

    # Save
    path = os.path.join(OUT, 'PEM en TIC_Sesion de Clase 6 Programacion I.docx')
    doc.save(path)
    print(f'Sesion saved: {path}')

# ═══════════════════════════════════════════════════════════
# 2. GUIA DE APRENDIZAJE 6
# ═══════════════════════════════════════════════════════════
def create_guia():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Universidad Rafael Landivar')
    r.bold = True; r.font.size = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Facultad de Humanidades').font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Departamento de Educacion').font.size = Pt(11)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Guia de aprendizaje num. 6')
    r.bold = True; r.font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Semana 6 - Unidad 6')
    r.bold = True; r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Fecha de entrega: Viernes 27 de febrero de 2026')
    r.bold = True; r.font.size = Pt(11)

    doc.add_paragraph()

    # Info table
    info_table = doc.add_table(rows=2, cols=2)
    info_table.style = 'Table Grid'
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.rows[0].cells[0].text = 'Profesorados con Especialidad en TIC'
    info_table.rows[0].cells[1].text = 'Nombre del curso: Introduccion al Desarrollo Web con Java'
    info_table.rows[1].cells[0].text = 'Ano Psicopedagogico'
    info_table.rows[1].cells[1].text = 'Enfoque Pedagogico para Docentes'
    for row in info_table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)

    doc.add_paragraph()

    # ── Parte introductoria ──
    doc.add_heading('PARTE INTRODUCTORIA', level=1)

    p = doc.add_paragraph()
    r = p.add_run('"La seguridad no es un obstaculo: es la estructura que permite confiar en el sistema."')
    r.italic = True; r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.add_run('Contexto (conexion con Semana 5): ').bold = True
    p.add_run('Ya se construyo el CRUD funcional con rutas, controladores, formularios y guardado en H2. '
              'En Semana 6 nos enfocamos en proteger esos datos con login y roles, buscar informacion con consultas '
              'personalizadas y visualizarla con graficos. El objetivo es que el grupo domine: '
              'Login → Roles → Vista condicional → Consultas → Graficos.')

    doc.add_paragraph()

    # Learning outcomes table
    add_styled_table(doc,
        ['Aprendizajes esperados', 'Productos que evidencian el aprendizaje'],
        [
            ['Implementa login con Spring Security y usuarios en memoria.',
             'Login funcional: evidencia de ingreso exitoso con admin y docente.'],
            ['Configura roles y protege rutas por permisos.',
             'Vista condicional: captura de ADMIN vs USER (diferente interfaz).'],
            ['Busca y filtra datos con metodos JPA sin escribir SQL.',
             'Filtro funcional: captura de busqueda por tipo y por texto.'],
            ['Visualiza datos de la BD en graficos interactivos.',
             'Grafico con datos reales: captura de barras/dona con conteo real.'],
            ['Explica el flujo de seguridad con analogias claras.',
             'Bitacora didactica: 3 analogias + explicacion del flujo en 10 lineas.']
        ])

    doc.add_paragraph()

    # ── PRIMERA PARTE ──
    doc.add_heading('PRIMERA PARTE: SEGURIDAD Y LOGIN (COMPRENDER ANTES DE CODIFICAR)', level=1)

    p = doc.add_paragraph()
    p.add_run('Proposito de la actividad: ').bold = True
    p.add_run('entender por que se protegen las rutas y como funciona el flujo de login antes de implementarlo.')

    doc.add_heading('Actividades (Instrucciones):', level=2)

    doc.add_heading('1) Lectura guiada (25-35 min)', level=3)
    doc.add_paragraph('- Abrir el material HTML de la semana: Material_Sesion_Clase_6_PERSISTENCIA/index.html')
    doc.add_paragraph('- Leer las secciones "Login paso a paso" y "Roles y permisos" y responder:')
    doc.add_paragraph('  - Que dependencia habilita el login en Spring Boot.')
    doc.add_paragraph('  - Cual es la diferencia entre ADMIN y USER en este sistema.')
    doc.add_paragraph('  - Por que las contrasenas se cifran (nunca texto plano).')
    doc.add_paragraph('  - Que pasa si un usuario sin login intenta acceder a /recursos.')

    doc.add_heading('2) Mapa de seguridad + roles (20-25 min)', level=3)
    doc.add_paragraph('- Crear un diagrama sencillo con el flujo de seguridad:')
    doc.add_paragraph('  - Usuario → Login → Spring Security verifica → Redirige segun resultado')
    doc.add_paragraph('  - ADMIN → ve todo (CRUD completo + estadisticas)')
    doc.add_paragraph('  - USER → ve listado + crear (sin borrar ni estadisticas)')
    doc.add_paragraph('  - Anonimo → redirige a /login')
    doc.add_paragraph('- Escribir 1 frase por rol: que puede hacer y que no.')

    p = doc.add_paragraph()
    p.add_run('Recursos: ').bold = True
    p.add_run('Material HTML Sesion 6 y notas de la Semana 5.')
    p = doc.add_paragraph()
    p.add_run('Evaluacion (Formativa): ').bold = True
    p.add_run('claridad del mapa y explicacion oral.')

    doc.add_paragraph()

    # ── SEGUNDA PARTE ──
    doc.add_heading('SEGUNDA PARTE: IMPLEMENTACION (PASO A PASO)', level=1)

    p = doc.add_paragraph()
    p.add_run('Proposito de la actividad: ').bold = True
    p.add_run('construir login funcional, renderizado condicional, filtros y graficos. La meta no es solo que funcione, sino poder explicarlo.')

    p = doc.add_paragraph()
    p.add_run('Entregable principal (equipo, 3-4 integrantes): ').bold = True
    p.add_run('Login funcional + vista condicional + filtro + grafico con datos reales')

    doc.add_heading('1) Checklist previo (equipo):', level=3)
    doc.add_paragraph('- Verificar que el CRUD de Semana 5 compila y funciona.')
    doc.add_paragraph('- Confirmar que hay al menos 5 registros en la BD.')
    doc.add_paragraph('- Verificar acceso a la consola H2.')

    doc.add_heading('2) Agregar dependencias (equipo):', level=3)
    doc.add_paragraph('- En pom.xml agregar: spring-boot-starter-security')
    doc.add_paragraph('- En pom.xml agregar: thymeleaf-extras-springsecurity6')
    doc.add_paragraph('- Ejecutar ./mvnw clean y verificar que compila.')

    doc.add_heading('3) Crear configuracion de seguridad (equipo):', level=3)
    doc.add_paragraph('- Crear carpeta configuracion/ en el proyecto.')
    doc.add_paragraph('- Crear SeguridadConfig.java con @Configuration + @EnableWebSecurity.')
    doc.add_paragraph('- Definir SecurityFilterChain con reglas de acceso.')
    doc.add_paragraph('- Definir UserDetailsService con admin y docente.')
    doc.add_paragraph('- Definir PasswordEncoder con BCryptPasswordEncoder.')

    doc.add_heading('4) Crear login (equipo):', level=3)
    doc.add_paragraph('- Crear LoginControlador.java con ruta GET /login.')
    doc.add_paragraph('- Crear templates/login.html con formulario (campos: username, password).')
    doc.add_paragraph('- Probar: login exitoso, login fallido, cerrar sesion.')

    doc.add_heading('5) Renderizado condicional (equipo):', level=3)
    doc.add_paragraph('- Agregar xmlns:sec al HTML de lista.html.')
    doc.add_paragraph('- Mostrar nombre del usuario conectado.')
    doc.add_paragraph('- Ocultar "Borrar" para USER; mostrar para ADMIN.')
    doc.add_paragraph('- Agregar enlace a "Estadisticas" solo para ADMIN.')
    doc.add_paragraph('- Agregar boton "Cerrar sesion" con formulario POST a /logout.')

    doc.add_heading('6) Consultas y filtros (equipo):', level=3)
    doc.add_paragraph('- Agregar findByTipo, findByTituloContainingIgnoreCase, countByTipo al repositorio.')
    doc.add_paragraph('- Modificar controlador para aceptar @RequestParam.')
    doc.add_paragraph('- Agregar formulario de busqueda en lista.html.')

    doc.add_heading('7) Graficos (equipo):', level=3)
    doc.add_paragraph('- Crear ruta GET /recursos/estadisticas.')
    doc.add_paragraph('- Crear estadisticas.html con Chart.js.')
    doc.add_paragraph('- Mostrar grafico de barras con conteo por tipo.')

    doc.add_paragraph()

    # Rubrica
    doc.add_heading('Evaluacion (rubrica rapida 100 pts):', level=2)
    add_styled_table(doc,
        ['Criterio', 'Puntos'],
        [
            ['Login funciona (admin + docente + error + logout)', '20'],
            ['Renderizado condicional correcto (ADMIN vs USER)', '20'],
            ['Filtro por tipo y busqueda funcionan', '20'],
            ['Grafico con datos reales de la BD', '20'],
            ['Explicacion + analogias + orden del codigo', '20']
        ])

    doc.add_paragraph()

    # ── CRONOGRAMA ──
    doc.add_heading('CRONOGRAMA (10 horas)', level=1)
    add_styled_table(doc,
        ['Actividad', 'Tiempo', 'Aprendizaje'],
        [
            ['Lectura guiada + mapa de seguridad', '1.5 horas', 'Comprension de login y roles'],
            ['Configuracion de seguridad + login', '2.5 horas', 'Spring Security funcional'],
            ['Renderizado condicional', '1.5 horas', 'Vista diferente por rol'],
            ['Consultas JPA y filtrado', '2.0 horas', 'Busquedas sin SQL'],
            ['Graficos con Chart.js', '1.5 horas', 'Visualizacion de datos'],
            ['Bitacora + reflexion docente', '1.0 horas', 'Metacognicion'],
            ['', 'Total: 10 horas', '']
        ])

    doc.add_paragraph()

    # ── RECURSOS ──
    doc.add_heading('RECURSOS DE APOYO', level=1)
    doc.add_paragraph('- Material HTML Sesion 6: Material_Sesion_Clase_6_PERSISTENCIA/')
    doc.add_paragraph('- Material Sesion 5 (rutas y CRUD): Material_Sesion_Clase_5_RUTAS/')
    doc.add_paragraph('- Spring Security (referencia): https://spring.io/projects/spring-security')
    doc.add_paragraph('- Chart.js (documentacion): https://www.chartjs.org/')
    doc.add_paragraph('- Spring Initializr: https://start.spring.io/')

    # Save
    path = os.path.join(OUT, 'PEM en TIC_Guia de Aprendizaje 6 Programacion 1.docx')
    doc.save(path)
    print(f'Guia saved: {path}')

if __name__ == '__main__':
    create_sesion()
    create_guia()
    print('Done!')
